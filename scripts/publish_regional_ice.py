"""Publish NSIDC regional workbook values and its matching 2007 mask boundaries."""
from __future__ import annotations
import argparse, calendar, hashlib, json, math, sys
from datetime import date, datetime, timezone
from pathlib import Path
import numpy as np
import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.analysis-deps'))
REGIONS = [('barents','巴伦支海','Barents',8,[73,38]),('kara','喀拉海','Kara',9,[75,77]),('chukchi','楚科奇海','Chukchi',12,[70,-170])]
BASE='https://noaadata.apps.nsidc.org/NOAA/G02135/seaice_analysis/'
BOOK='N_Sea_Ice_Index_Regional_Monthly_Data_G02135_v4.0.xlsx'
MASK='Arctic_region_mask_Meier_AnnGlaciol2007.msk'
OUT=ROOT/'data/analysis'

def atomic(path, value):
    temp=path.with_suffix(path.suffix+'.tmp')
    temp.write_text(json.dumps(value,ensure_ascii=False,allow_nan=False,separators=(',',':')),encoding='utf-8')
    temp.replace(path)

def asset(name, refresh=False):
    path=OUT/'raw'/name;path.parent.mkdir(parents=True,exist_ok=True)
    # Preserve prior verified source bytes before any refresh, independently of published outputs.
    if path.exists() and path.with_name(name+'.source.json').exists():
        old=json.loads(path.with_name(name+'.source.json').read_text(encoding='utf-8'));body=path.read_bytes()
        if hashlib.sha256(body).hexdigest()!=old['sha256']:raise ValueError('Existing raw hash mismatch')
        archive=OUT/'raw/versions'/old['sha256'];archive.mkdir(parents=True,exist_ok=True)
        if not (archive/name).exists():(archive/name).write_bytes(body)
        if not (archive/(name+'.source.json')).exists():atomic(archive/(name+'.source.json'),old)
    if refresh or not path.exists():
        response=requests.get(BASE+name,timeout=45);response.raise_for_status()
        if not 1000<len(response.content)<5_000_000:raise ValueError('Unexpected source length')
        path.write_bytes(response.content)
        atomic(path.with_name(name+'.source.json'),dict(url=BASE+name,sha256=hashlib.sha256(response.content).hexdigest(),bytes=len(response.content),retrieved_at=datetime.now(timezone.utc).isoformat()))
    meta=json.loads(path.with_name(name+'.source.json').read_text(encoding='utf-8'))
    if hashlib.sha256(path.read_bytes()).hexdigest()!=meta['sha256']:raise ValueError('Raw hash mismatch')
    return path,meta

def parse_workbook(path, source):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    rows={};audit=[]
    for rid,name,sheet,code,center in REGIONS:
        for metric in ['area','extent']:
            sheet_name=f'{sheet}-{metric.title()}-km^2';ws=wb[sheet_name]
            values=list(ws.values)
            assert len(values[0])==25 and values[0][1::2]==tuple(calendar.month_name[1:]),sheet_name
            assert all(v==metric for v in values[1][1::2]),sheet_name
            seen=set();numeric=0
            for line,record in enumerate(values[3:],4):
                year=record[0]
                assert isinstance(year,int) and 1978<=year<=date.today().year and year not in seen,(sheet_name,line)
                seen.add(year)
                for month in range(1,13):
                    raw=record[month*2-1]
                    if raw is not None:
                        assert isinstance(raw,(int,float)) and math.isfinite(raw) and 0<=raw<5_000_000,(sheet_name,line,month)
                        assert (year,month)<(date.today().year,date.today().month),'Only finished months are published'
                        numeric+=1
                    key=(rid,year,month)
                    row=rows.setdefault(key,dict(region_id=rid,year=year,month=month,date=f'{year:04}-{month:02}',source_id='nsidc-regional:'+source['sha256']))
                    row[metric+'_km2']=raw
                    row[metric+'_cell']=sheet_name+'!'+openpyxl.utils.get_column_letter(month*2)+str(line)
            audit.append(dict(sheet=sheet_name,raw_rows=ws.max_row,data_years=len(seen),first_year=min(seen),last_year=max(seen),numeric_months=numeric))
    observed=[r['date'] for r in rows.values() if r.get('extent_km2') is not None]
    first,last=min(observed),max(observed)
    result=sorted((r for r in rows.values() if first<=r['date']<=last),key=lambda r:(r['region_id'],r['date']))
    for r in result:
        assert (r['area_km2'] is None)==(r['extent_km2'] is None),'Mismatched area/extent coverage'
        if r['area_km2'] is not None:assert r['area_km2']<=r['extent_km2']+.01
        r['status']='observed' if r['extent_km2'] is not None else 'missing'
        for metric in ['area','extent']:
            base=[x[metric+'_km2'] for x in result if x['region_id']==r['region_id'] and x['month']==r['month'] and 1991<=x['year']<=2020 and x[metric+'_km2'] is not None]
            r[metric+'_baseline_n']=len(base)
            r[metric+'_baseline_km2']=round(sum(base)/30,3) if len(base)==30 else None
            r[metric+'_anomaly_km2']=round(r[metric+'_km2']-r[metric+'_baseline_km2'],3) if r[metric+'_km2'] is not None and len(base)==30 else None
    wb.close()
    return result,audit

def boundaries(path):
    from pyproj import Transformer
    from shapely import box, union_all, Polygon, MultiPolygon
    from shapely.geometry import mapping
    from shapely.affinity import translate
    mask=np.frombuffer(path.read_bytes(),dtype=np.uint8).reshape(448,304)
    transformer=Transformer.from_crs('EPSG:3411','EPSG:4326',always_xy=True)
    # Grid outer corner and spacing from NSIDC's polar stereographic grid guide.
    assert np.allclose(transformer.transform(-3850000,5850000),(168.35,30.98),atol=.02)
    def ring(coords):
        result=[]
        for x,y in coords:
            lon,lat=transformer.transform(x,y)
            if result:
                while lon-result[-1][0]>180:lon-=360
                while lon-result[-1][0]<-180:lon+=360
            result.append((lon,lat))
        return result
    features=[]
    for rid,name,sheet,code,center in REGIONS:
        ys,xs=np.where(mask==code)
        cells=[box(-3850000+x*25000,5850000-(y+1)*25000,-3850000+(x+1)*25000,5850000-y*25000) for y,x in zip(ys,xs)]
        merged=union_all(cells)
        polys=list(merged.geoms) if merged.geom_type=='MultiPolygon' else [merged]
        pieces=[]
        for poly in polys:
            exterior=ring(poly.exterior.coords);origin=np.mean([x for x,y in exterior])
            holes=[]
            for interior in poly.interiors:
                h=ring(interior.coords);offset=round((origin-np.mean([x for x,y in h]))/360)*360
                holes.append([(x+offset,y) for x,y in h])
            geographic=Polygon(exterior,holes)
            assert geographic.is_valid,rid
            for shift in [-360,0,360]:
                clip=geographic.intersection(box(-180+shift,-90,180+shift,90))
                if clip.is_empty:continue
                for part in ([clip] if clip.geom_type=='Polygon' else list(clip.geoms)):
                    if part.geom_type=='Polygon':pieces.append(translate(part,xoff=-shift))
        geo=MultiPolygon(pieces)
        features.append(dict(type='Feature',properties=dict(id=rid,name=name,mask_value=code,grid_cells=len(cells),definition='NSIDC Meier et al. 2007 sea-ice region mask; 25 km cells'),geometry=mapping(geo)))
    return dict(type='FeatureCollection',features=features)

def main():
    if not __debug__:raise RuntimeError('Run without Python -O: source validation must remain enabled')
    args=argparse.ArgumentParser();args.add_argument('--refresh',action='store_true');opt=args.parse_args()
    book,book_meta=asset(BOOK,opt.refresh);mask,mask_meta=asset(MASK,opt.refresh)
    records,audit=parse_workbook(book,book_meta);geo=boundaries(mask)
    definitions=[dict(id=r[0],name=r[1],sheet_prefix=r[2],mask_value=r[3],center=r[4]) for r in REGIONS]
    result=dict(version='regional-v1-'+book_meta['sha256'][:12],unit='km²',baseline=[1991,2020],historical_window=[2016,2025],regions=definitions,records=records,
                sources=[dict(id='nsidc-regional:'+book_meta['sha256'],title='NSIDC Sea Ice Index V4 · 北极分区月度数据',**book_meta),dict(id='nsidc-region-mask:'+mask_meta['sha256'],title='NSIDC Meier 2007 海区掩膜',**mask_meta)],
                methods=dict(reference='https://nsidc.org/sites/default/files/documents/technical-reference/sea-ice-analysis-spreadsheets-overview.pdf',grid_reference='https://nsidc.org/ru/node/52236',credit='Credit: Sea Ice Index, National Snow and Ice Data Center.',spatial='海冰指标直接读取官方分区表；地图轮廓由配套 25 km 掩膜转换，不用于领土或法律边界判断。',comparison='相同海区、相同月份；保留缺失，不把面积与范围混用；1991—2020 同月基准须有 30 个有效年份。',navigation='月度海冰总量不换算为航道通航天数。'),workbook_audit=audit)
    # Publish only after every sheet and geometry has passed validation.
    atomic(OUT/'regional_ice.json',result);atomic(OUT/'regions.geojson',geo)
    print(json.dumps(dict(regions=len(definitions),rows=len(records),first=min(r['date'] for r in records),last=max(r['date'] for r in records),missing=sum(r['status']=='missing' for r in records)),ensure_ascii=False))

if __name__=='__main__':main()
